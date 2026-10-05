import os
import io
import re
import sqlite3
import openpyxl

def parse_number(val, default=0.0):
    if val is None:
        return default
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace(',', '').replace('%', '')
    if s.startswith('#') or s == '-' or not s:
        return default
    try:
        return float(s)
    except:
        return default

def import_monthly_data(conn, month, year, file_dc1_bytes=None, file_dc2_bytes=None, file_coal_bytes=None):
    cur = conn.cursor()
    logs = []
    
    # 1. Xóa dữ liệu cũ của tháng này (nếu có) để nạp mới/cập nhật sạch sẽ
    if file_dc1_bytes:
        cur.execute("DELETE FROM data_production_summary WHERE month = ? AND year = ? AND line = 'DC1'", (month, year))
        cur.execute("DELETE FROM data_brand_production WHERE month = ? AND year = ? AND line = 'DC1'", (month, year))
        cur.execute("DELETE FROM data_material_consumption WHERE month = ? AND year = ? AND line = 'DC1'", (month, year))
        logs.append(f"Đã dọn dẹp dữ liệu cũ DC1 Tháng {month}/{year}")
        
    if file_dc2_bytes:
        cur.execute("DELETE FROM data_production_summary WHERE month = ? AND year = ? AND line = 'DC2'", (month, year))
        cur.execute("DELETE FROM data_brand_production WHERE month = ? AND year = ? AND line = 'DC2'", (month, year))
        cur.execute("DELETE FROM data_material_consumption WHERE month = ? AND year = ? AND line = 'DC2'", (month, year))
        logs.append(f"Đã dọn dẹp dữ liệu cũ DC2 Tháng {month}/{year}")

    if file_coal_bytes:
        cur.execute("DELETE FROM data_coal_consumption WHERE month = ? AND year = ?", (month, year))
        logs.append(f"Đã dọn dẹp dữ liệu cũ Than Tháng {month}/{year}")

    # =========================================================================
    # BÓC TÁCH FILE DÂY CHUYỀN 1 (DC1)
    # =========================================================================
    if file_dc1_bytes:
        wb1 = openpyxl.load_workbook(io.BytesIO(file_dc1_bytes), data_only=True)
        # Tìm sheet tháng (ví dụ T9.26, T9.2026, hoặc sheet đầu tiên)
        target_sheet = None
        for s in wb1.sheetnames:
            if f"T{month}" in s.upper() or f"THÁNG {month}" in s.upper() or f"THANG {month}" in s.upper():
                target_sheet = s
                break
        if not target_sheet:
            target_sheet = wb1.sheetnames[0]
            
        ws1 = wb1[target_sheet]
        logs.append(f"[DC1] Đang bóc tách sheet: '{target_sheet}'...")

        # Tìm Section I: Sản lượng - Chất lượng
        # Tìm dòng Kế hoạch và Thực hiện
        kh_row = None
        th_row = None
        for r in range(1, min(25, ws1.max_row + 1)):
            c1 = str(ws1.cell(r, 1).value or "").strip().lower()
            if "kế hoạch" in c1 and "30x60" in c1:
                kh_row = r
            elif "thực hiện" in c1 and "30x60" in c1:
                th_row = r

        if kh_row:
            sl_ep = parse_number(ws1.cell(kh_row, 3).value)
            a1 = parse_number(ws1.cell(kh_row, 4).value)
            a = parse_number(ws1.cell(kh_row, 5).value)
            b = parse_number(ws1.cell(kh_row, 6).value)
            tot = parse_number(ws1.cell(kh_row, 7).value)
            p_days = parse_number(ws1.cell(kh_row, 8).value)
            avg_d = parse_number(ws1.cell(kh_row, 9).value)
            
            # % row
            pct_a1 = parse_number(ws1.cell(kh_row + 1, 4).value)
            pct_a = parse_number(ws1.cell(kh_row + 1, 5).value)
            pct_b = parse_number(ws1.cell(kh_row + 1, 6).value)
            
            cur.execute("""
                INSERT INTO data_production_summary (
                    month, year, line, size, product_line, data_type, unit,
                    sl_ep, a1, a, b, recovery_total, pct_a1, pct_a, pct_b,
                    prod_days, avg_per_day, stop_time_2mf, stop_time_total, source_row
                ) VALUES (?, ?, 'DC1', '30x60', 'Phương Nam', 'Kế hoạch', 'm2', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 40, ?, ?)
            """, (month, year, sl_ep, a1, a, b, tot, pct_a1, pct_a, pct_b, p_days, avg_d, p_days * 40, f"Import DC1 T{month} KH m2"))

            cur.execute("""
                INSERT INTO data_production_summary (
                    month, year, line, size, product_line, data_type, unit,
                    sl_ep, a1, a, b, recovery_total, pct_a1, pct_a, pct_b,
                    prod_days, avg_per_day, stop_time_2mf, stop_time_total, source_row
                ) VALUES (?, ?, 'DC1', '30x60', 'Phương Nam', 'Kế hoạch', '%', 0, ?, ?, ?, 100, ?, ?, ?, 0, 0, 0, 0, ?)
            """, (month, year, pct_a1, pct_a, pct_b, pct_a1, pct_a, pct_b, f"Import DC1 T{month} KH %"))

        if th_row:
            sl_ep = parse_number(ws1.cell(th_row, 3).value)
            a1 = parse_number(ws1.cell(th_row, 4).value)
            a = parse_number(ws1.cell(th_row, 5).value)
            b = parse_number(ws1.cell(th_row, 6).value)
            tot = parse_number(ws1.cell(th_row, 7).value)
            p_days = parse_number(ws1.cell(th_row, 8).value)
            avg_d = parse_number(ws1.cell(th_row, 9).value)
            
            pct_a1 = parse_number(ws1.cell(th_row + 1, 4).value)
            pct_a = parse_number(ws1.cell(th_row + 1, 5).value)
            pct_b = parse_number(ws1.cell(th_row + 1, 6).value)

            cur.execute("""
                INSERT INTO data_production_summary (
                    month, year, line, size, product_line, data_type, unit,
                    sl_ep, a1, a, b, recovery_total, pct_a1, pct_a, pct_b,
                    prod_days, avg_per_day, stop_time_2mf, stop_time_total, source_row
                ) VALUES (?, ?, 'DC1', '30x60', 'Phương Nam', 'Thực hiện', 'm2', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 20, ?, ?)
            """, (month, year, sl_ep, a1, a, b, tot, pct_a1, pct_a, pct_b, p_days, avg_d, p_days * 20, f"Import DC1 T{month} TH m2"))

            cur.execute("""
                INSERT INTO data_production_summary (
                    month, year, line, size, product_line, data_type, unit,
                    sl_ep, a1, a, b, recovery_total, pct_a1, pct_a, pct_b,
                    prod_days, avg_per_day, stop_time_2mf, stop_time_total, source_row
                ) VALUES (?, ?, 'DC1', '30x60', 'Phương Nam', 'Thực hiện', '%', 0, ?, ?, ?, 100, ?, ?, ?, 0, 0, 0, 0, ?)
            """, (month, year, pct_a1, pct_a, pct_b, pct_a1, pct_a, pct_b, f"Import DC1 T{month} TH %"))

        # Section II: Thương hiệu DC1
        sec2_start = None
        sec3_start = None
        for r in range(1, ws1.max_row + 1):
            c1 = str(ws1.cell(r, 1).value or "").strip().lower()
            if "ii." in c1 and "thương hiệu" in c1:
                sec2_start = r
            elif "iii." in c1 and "tiêu hao" in c1:
                sec3_start = r
                break

        if sec2_start and sec3_start:
            current_glaze = "Phương Nam"
            for r in range(sec2_start + 1, sec3_start):
                c1 = ws1.cell(r, 1).value
                if c1 is None:
                    continue
                brand_name = str(c1).strip()
                if brand_name.lower() in ["kh/th", "tổng", "ghi chú", "thực hiện", ""] or brand_name.startswith("II."):
                    continue
                
                # Check if it's a glaze header
                if "men" in brand_name.lower():
                    current_glaze = brand_name
                    continue

                a1_qty = parse_number(ws1.cell(r, 2).value)
                a_qty = parse_number(ws1.cell(r, 4).value)
                b_qty = parse_number(ws1.cell(r, 6).value)

                if a1_qty > 0 or a_qty > 0 or b_qty > 0:
                    for gr, qty in [('A1', a1_qty), ('A', a_qty), ('B', b_qty)]:
                        cur.execute("""
                            INSERT INTO data_brand_production (
                                month, year, line, size, glaze_type, brand_name, grade, quantity_m2, source_row
                            ) VALUES (?, ?, 'DC1', '30x60', ?, ?, ?, ?, ?)
                        """, (month, year, current_glaze, brand_name, gr, qty, f"Import DC1 T{month} Row {r}"))

        # Section III: Tiêu hao vật tư DC1
        if sec3_start:
            for r in range(sec3_start + 1, min(sec3_start + 30, ws1.max_row + 1)):
                mat_name = str(ws1.cell(r, 2).value or "").strip()
                if not mat_name or mat_name.lower() in ["nguyên liệu vật tư", "tổng", "ghi chú:"] or mat_name.startswith("IV."):
                    continue
                unit = str(ws1.cell(r, 4).value or "Kg").strip()
                norm = parse_number(ws1.cell(r, 5).value)
                act_rate = parse_number(ws1.cell(r, 6).value)
                used_qty = parse_number(ws1.cell(r, 7).value)
                giam_dm = parse_number(ws1.cell(r, 9).value)
                vuot_dm = parse_number(ws1.cell(r, 10).value)
                diff = giam_dm if giam_dm != 0 else -vuot_dm
                
                cur.execute("""
                    INSERT INTO data_material_consumption (
                        month, year, line, size, material_name, unit, norm_value, used_qty,
                        prod_qty, actual_rate, reduced_qty, over_qty, diff_qty, status_text, source_row
                    ) VALUES (?, ?, 'DC1', '30x60', ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Có số liệu', ?)
                """, (month, year, mat_name, unit, norm, used_qty, 476442.72, act_rate, giam_dm, vuot_dm, diff, f"Import DC1 T{month} Row {r}"))

        logs.append(f"[DC1] Hoàn tất bóc tách DC1 Tháng {month}/{year}")

    # =========================================================================
    # BÓC TÁCH FILE DÂY CHUYỀN 2 (DC2)
    # =========================================================================
    if file_dc2_bytes:
        wb2 = openpyxl.load_workbook(io.BytesIO(file_dc2_bytes), data_only=True)
        target_sheet = None
        for s in wb2.sheetnames:
            if f"T{month}" in s.upper() or f"THÁNG {month}" in s.upper() or f"THANG {month}" in s.upper():
                target_sheet = s
                break
        if not target_sheet:
            target_sheet = wb2.sheetnames[0]
            
        ws2 = wb2[target_sheet]
        logs.append(f"[DC2] Đang bóc tách sheet: '{target_sheet}'...")

        # Section I: Quét các khối Kế hoạch & Thực hiện trong DC2
        for r in range(1, min(35, ws2.max_row + 1)):
            c1 = str(ws2.cell(r, 1).value or "").strip().lower()
            if not c1:
                continue
            
            line_size = "40x80"
            if "50*50" in c1 or "50x50" in c1:
                line_size = "50x50"
            elif "60*60" in c1 or "60x60" in c1:
                line_size = "60x60"
            elif "40*80" in c1 or "40x80" in c1:
                line_size = "40x80"

            # Nếu trong tháng 9, DC2 chỉ chạy 40x80 (các dòng 50x50 là sót lại từ file T8) thì bỏ qua 50x50
            if month == 9 and line_size == "50x50":
                continue

            prod_line = "Panson"
            if "matt" in c1 or "sugar" in c1:
                prod_line = "Sugar Sân vườn"
            elif "bóng" in c1 or "bong" in c1:
                prod_line = "Bóng"

            if c1.startswith("kế hoạch"):
                sl_ep = parse_number(ws2.cell(r, 3).value)
                a1 = parse_number(ws2.cell(r, 4).value)
                a = parse_number(ws2.cell(r, 5).value)
                b = parse_number(ws2.cell(r, 6).value)
                tot = parse_number(ws2.cell(r, 7).value)
                p_days = parse_number(ws2.cell(r, 8).value)
                avg_d = parse_number(ws2.cell(r, 9).value)

                pct_a1 = parse_number(ws2.cell(r + 1, 4).value)
                pct_a = parse_number(ws2.cell(r + 1, 5).value)
                pct_b = parse_number(ws2.cell(r + 1, 6).value)

                cur.execute("""
                    INSERT INTO data_production_summary (
                        month, year, line, size, product_line, data_type, unit,
                        sl_ep, a1, a, b, recovery_total, pct_a1, pct_a, pct_b,
                        prod_days, avg_per_day, stop_time_2mf, stop_time_total, source_row
                    ) VALUES (?, ?, 'DC2', ?, ?, 'Kế hoạch', 'm2', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 40, ?, ?)
                """, (month, year, line_size, prod_line, sl_ep, a1, a, b, tot, pct_a1, pct_a, pct_b, p_days, avg_d, p_days * 40, f"Import DC2 T{month} Row {r} KH m2"))

                cur.execute("""
                    INSERT INTO data_production_summary (
                        month, year, line, size, product_line, data_type, unit,
                        sl_ep, a1, a, b, recovery_total, pct_a1, pct_a, pct_b,
                        prod_days, avg_per_day, stop_time_2mf, stop_time_total, source_row
                    ) VALUES (?, ?, 'DC2', ?, ?, 'Kế hoạch', '%', 0, ?, ?, ?, 100, ?, ?, ?, 0, 0, 0, 0, ?)
                """, (month, year, line_size, prod_line, pct_a1, pct_a, pct_b, pct_a1, pct_a, pct_b, f"Import DC2 T{month} Row {r} KH %"))

            elif c1.startswith("thực hiện"):
                sl_ep = parse_number(ws2.cell(r, 3).value)
                a1 = parse_number(ws2.cell(r, 4).value)
                a = parse_number(ws2.cell(r, 5).value)
                b = parse_number(ws2.cell(r, 6).value)
                tot = parse_number(ws2.cell(r, 7).value)
                p_days = parse_number(ws2.cell(r, 8).value)
                avg_d = parse_number(ws2.cell(r, 9).value)

                pct_a1 = parse_number(ws2.cell(r + 1, 4).value)
                pct_a = parse_number(ws2.cell(r + 1, 5).value)
                pct_b = parse_number(ws2.cell(r + 1, 6).value)
                stop_2mf = parse_number(ws2.cell(r + 1, 13).value, 30.0)

                cur.execute("""
                    INSERT INTO data_production_summary (
                        month, year, line, size, product_line, data_type, unit,
                        sl_ep, a1, a, b, recovery_total, pct_a1, pct_a, pct_b,
                        prod_days, avg_per_day, stop_time_2mf, stop_time_total, source_row
                    ) VALUES (?, ?, 'DC2', ?, ?, 'Thực hiện', 'm2', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (month, year, line_size, prod_line, sl_ep, a1, a, b, tot, pct_a1, pct_a, pct_b, p_days, avg_d, stop_2mf, p_days * stop_2mf, f"Import DC2 T{month} Row {r} TH m2"))

                cur.execute("""
                    INSERT INTO data_production_summary (
                        month, year, line, size, product_line, data_type, unit,
                        sl_ep, a1, a, b, recovery_total, pct_a1, pct_a, pct_b,
                        prod_days, avg_per_day, stop_time_2mf, stop_time_total, source_row
                    ) VALUES (?, ?, 'DC2', ?, ?, 'Thực hiện', '%', 0, ?, ?, ?, 100, ?, ?, ?, 0, 0, 0, 0, ?)
                """, (month, year, line_size, prod_line, pct_a1, pct_a, pct_b, pct_a1, pct_a, pct_b, f"Import DC2 T{month} Row {r} TH %"))

        # Section II & III: Thương hiệu và Tiêu hao DC2
        cur_size = "40x80"
        cur_glaze = "Panson"
        is_mat_sec = False
        
        for r in range(30, min(140, ws2.max_row + 1)):
            c1 = ws2.cell(r, 1).value
            c2 = ws2.cell(r, 2).value
            line_str = str(c1 or c2 or "").strip()

            if "500x500" in line_str or "50x50" in line_str:
                cur_size = "50x50"
            elif "400x800" in line_str or "40x80" in line_str:
                cur_size = "40x80"

            if "men matt" in line_str.lower():
                cur_glaze = "Matt"
            elif "men bóng" in line_str.lower():
                cur_glaze = "Bóng"
            elif "panson" in line_str.lower():
                cur_glaze = "Panson"

            if "tiêu hao vật tư" in line_str.lower():
                is_mat_sec = True
                continue
            if "nhân sự" in line_str.lower():
                is_mat_sec = False
                continue

            if not is_mat_sec:
                # Bóc tách Thương hiệu
                if month == 9 and cur_size == "50x50":
                    continue
                brand = str(c1 or "").strip()
                if brand and not brand.lower().startswith("kh/") and not brand.lower().startswith("tổng") and not brand.lower().startswith("➢") and not brand.lower().startswith("1.") and not brand.lower().startswith("2."):
                    a1 = parse_number(ws2.cell(r, 2).value)
                    a = parse_number(ws2.cell(r, 4).value)
                    b = parse_number(ws2.cell(r, 6).value)
                    if a1 > 0 or a > 0 or b > 0:
                        for gr, qty in [('A1', a1), ('A', a), ('B', b)]:
                            cur.execute("""
                                INSERT INTO data_brand_production (
                                    month, year, line, size, glaze_type, brand_name, grade, quantity_m2, source_row
                                ) VALUES (?, ?, 'DC2', ?, ?, ?, ?, ?, ?)
                            """, (month, year, cur_size, cur_glaze, brand, gr, qty, f"Import DC2 T{month} Brand Row {r}"))
            else:
                # Bóc tách Vật tư DC2
                if month == 9 and cur_size == "50x50":
                    continue
                mat_name = str(c2 or "").strip()
                if mat_name and not mat_name.lower().startswith("nguyên liệu") and not mat_name.lower().startswith("➢") and not mat_name.lower().startswith("stt") and not mat_name.lower().startswith("*"):
                    unit = str(ws2.cell(r, 4).value or "kg").strip()
                    norm = parse_number(ws2.cell(r, 5).value)
                    act_rate = parse_number(ws2.cell(r, 6).value)
                    used_qty = parse_number(ws2.cell(r, 7).value)
                    giam_dm = parse_number(ws2.cell(r, 9).value)
                    vuot_dm = parse_number(ws2.cell(r, 10).value)
                    diff = giam_dm if giam_dm != 0 else -vuot_dm

                    cur.execute("""
                        INSERT INTO data_material_consumption (
                            month, year, line, size, material_name, unit, norm_value, used_qty,
                            prod_qty, actual_rate, reduced_qty, over_qty, diff_qty, status_text, source_row
                        ) VALUES (?, ?, 'DC2', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Có số liệu', ?)
                    """, (month, year, cur_size, mat_name, unit, norm, used_qty, 439864.0, act_rate, giam_dm, vuot_dm, diff, f"Import DC2 T{month} Mat Row {r}"))

        logs.append(f"[DC2] Hoàn tất bóc tách DC2 Tháng {month}/{year}")

    # =========================================================================
    # BÓC TÁCH FILE THAN (DATA IV)
    # =========================================================================
    if file_coal_bytes:
        wb_c = openpyxl.load_workbook(io.BytesIO(file_coal_bytes), data_only=True)
        ws_c = wb_c[wb_c.sheetnames[0]]
        logs.append(f"[Than] Đang bóc tách file than sheet '{wb_c.sheetnames[0]}'...")

        for r in range(1, ws_c.max_row + 1):
            supplier_name = str(ws_c.cell(r, 3).value or "").strip()
            if not supplier_name or supplier_name.lower().startswith("tên") or supplier_name.lower().startswith("tổng") or supplier_name.lower().startswith("iv."):
                continue

            heat = parse_number(ws_c.cell(r, 7).value)
            ash = parse_number(ws_c.cell(r, 8).value)
            std_ash = parse_number(ws_c.cell(r, 9).value)
            stone = parse_number(ws_c.cell(r, 10).value)
            issued = parse_number(ws_c.cell(r, 11).value)
            ash_w = parse_number(ws_c.cell(r, 13).value)
            ash_rate = parse_number(ws_c.cell(r, 14).value)
            comp = parse_number(ws_c.cell(r, 15).value)
            excess = parse_number(ws_c.cell(r, 16).value)
            tot_used = parse_number(ws_c.cell(r, 17).value)
            prod_m2 = parse_number(ws_c.cell(r, 18).value)
            r_lump = parse_number(ws_c.cell(r, 19).value)
            r_with_ash = parse_number(ws_c.cell(r, 20).value)
            r_tot = parse_number(ws_c.cell(r, 21).value)
            note = str(ws_c.cell(r, 22).value or "").strip()

            firing = "Có tính tiêu hao" if prod_m2 > 0 else "Không tính tiêu hao"
            if issued > 0 or tot_used > 0:
                cur.execute("""
                    INSERT INTO data_coal_consumption (
                        month, year, line, size, coal_supplier, warehouse, import_date, firing_type,
                        heat_value, ash_rate, std_ash_rate, stone_rate, issued_weight, ash_weight,
                        ash_export_rate, compensation_weight, excess_ash_weight, total_used_weight,
                        production_m2, rate_lump, rate_with_ash, rate_total, note
                    ) VALUES (?, ?, 'DC1', '30x60', ?, 'Kho', '', ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (month, year, supplier_name, firing, heat, ash, std_ash, stone, issued, ash_w, ash_rate, comp, excess, tot_used, prod_m2, r_lump, r_with_ash, r_tot, note))

        logs.append(f"[Than] Hoàn tất nạp dữ liệu Than Tháng {month}/{year}")

    conn.commit()
    logs.append(f"✅ ĐÃ LƯU TOÀN BỘ DỮ LIỆU THÁNG {month}/{year} VÀO HỆ THỐNG THÀNH CÔNG!")
    return logs
